import pdfplumber
import docx

def extract_text_from_pdf(file):
    extracted_text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_text += text + "\n"
    return extracted_text

def extract_text_from_docx(file):
    document = docx.Document(file)
    extracted_text = ""

    for para in document.paragraphs:
        if para.text:
            extracted_text += para.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    extracted_text += cell.text + "\n"

    return extracted_text

def extract_text_from_txt(file):
    content = file.read()
    if isinstance(content, bytes):
        content = content.decode("utf-8", errors="ignore")
    return content

def extract_text(file, filename):
    ext = filename.split(".")[-1].lower()
    if ext == "pdf":
        return extract_text_from_pdf(file)
    elif ext in ("doc", "docx"):
        return extract_text_from_docx(file)
    elif ext == "txt":
        return extract_text_from_txt(file)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")